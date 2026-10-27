from __future__ import annotations

import json
from xml.sax.saxutils import escape

import re

import sys

import zipfile

from io import BytesIO

from pathlib import Path

from typing import Any

import pandas as pd

import streamlit as st

from datetime import datetime, timezone, timedelta

PROJECT_ROOT = Path(__file__).resolve().parents[1]

if str(PROJECT_ROOT) not in sys.path:

    sys.path.insert(0, str(PROJECT_ROOT))

from app.api_client import (  # noqa: E402

    create_review,

    get_review_by_id,

    get_reviews,

    update_review_assessment,

)

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable,
    LongTable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


st.set_page_config(page_title="MediReview", page_icon="💊", layout="wide")

# ============================================================

# CONSTANTS

# ============================================================

REVIEWER_OPTIONS = [

    "Needs Review",

    "Side Effect",

    "No Side Effect",

    "Mixed Evidence",

    "Insufficient Information",

]

FILTER_OPTIONS = ["All", *REVIEWER_OPTIONS[1:]]

VIEWS = ["New case", "Current case", "History"]

BADGE_COLORS = {

    "Side Effect": "red",

    "Mixed Evidence": "orange",

    "No Side Effect": "green",

    "Insufficient Information": "gray",

    "Needs Review": "blue",

}

# ============================================================

# STATE

# ============================================================

def initialize_state() -> None:

    defaults = {

        "active_view": "New case",

        "goto": None,

        "analysis_result": None,

        "submitted_text": "",

        "source_name": None,

        "review_id": None,

        "reviewer_classification": "Needs Review",

        "reviewer_notes": "",

        "history_reviews": [],

        "history_loaded": False,

        "history_error": None,

    }

    for key, value in defaults.items():

        st.session_state.setdefault(key, value)

def reset_current_case() -> None:

    """Start a new case. Saved data is never deleted."""

    for key in ("manual_review", "patient_reference_input", "report_reference_input", "uploaded_report"):

        st.session_state.pop(key, None)

    st.session_state.update(

        analysis_result=None,

        submitted_text="",

        source_name=None,

        review_id=None,

        reviewer_classification="Needs Review",

        reviewer_notes="",

        active_view="New case",

    )

def open_saved_case(review_id: int) -> None:

    """Callback: load a saved case into the workspace."""

    try:

        review = get_review_by_id(review_id)

        result = review.get("analysis_result")

        if not isinstance(result, dict):

            raise ValueError("Saved review has no valid analysis result.")

        result = dict(result)

        result["review_id"] = review.get("id")

        result["patient_reference"] = review.get("patient_reference")

        result["report_reference"] = review.get("report_reference")

        classification = review.get("reviewer_classification") or "Needs Review"

        if classification not in REVIEWER_OPTIONS:

            classification = "Needs Review"

        st.session_state.update(

            analysis_result=result,

            review_id=review.get("id"),

            submitted_text=review.get("text") or "",

            source_name=review.get("source_name") or "Saved review",

            reviewer_classification=classification,

            reviewer_notes=review.get("reviewer_notes") or "",

            active_view="Current case",

            history_error=None,

        )

        st.session_state.pop(f"reviewer_classification_select_{review_id}", None)

        st.session_state.pop(f"reviewer_notes_field_{review_id}", None)

    except Exception as exc:

        st.session_state["history_error"] = f"Unable to open case: {exc}"

# ============================================================

# HELPERS

# ============================================================

def safe_text(value: Any, fallback: str = "—") -> str:

    text = "" if value is None else str(value).strip()

    return text or fallback

def unique_strings(values: Any) -> list[str]:

    if not isinstance(values, (list, tuple)):

        return []

    seen: set[str] = set()

    out: list[str] = []

    for value in values:

        text = str(value).strip() if value is not None else ""

        if text and text.casefold() not in seen:

            seen.add(text.casefold())

            out.append(text)

    return out

def show_badge(value: Any) -> None:

    text = safe_text(value, "Not classified")

    color = BADGE_COLORS.get(text, "gray")

    if hasattr(st, "badge"):

        st.badge(text, color=color)

    else:

        st.markdown(f":{color}[**{text}**]")

def get_entity_values(result: dict[str, Any], entity_type: str) -> list[str]:

    entities = result.get("entities")

    if not isinstance(entities, list):

        return []

    return unique_strings([

        e.get("value")

        for e in entities

        if isinstance(e, dict) and str(e.get("type", "")).strip().upper() == entity_type

    ])

def _span_start(entity: dict[str, Any]) -> int:

    try:

        return int((entity.get("span") or {}).get("char_start", 0))

    except (TypeError, ValueError):

        return 0

def _new_med(name: str) -> dict[str, Any]:

    return {"medicine": name, "dosage": None, "frequency": None, "route": None}

def merge_duplicate_medications(details: list[Any]) -> list[dict[str, Any]]:

    merged: dict[str, dict[str, Any]] = {}

    for item in details:

        if not isinstance(item, dict):

            continue

        name = str(item.get("medicine") or item.get("medication") or item.get("name") or "").strip()

        if not name:

            continue

        current = merged.setdefault(name.casefold(), _new_med(name))

        for field, value in (

            ("dosage", item.get("dosage") or item.get("dose") or item.get("strength")),

            ("frequency", item.get("frequency")),

            ("route", item.get("route")),

        ):

            if value:

                current[field] = value

    return [

        item for item in merged.values()

        if item["medicine"].casefold() not in {"none", "null"}

    ]

def build_medications_from_entities(entities: list[Any]) -> list[dict[str, Any]]:

    medicines: dict[str, dict[str, Any]] = {}

    current: dict[str, Any] | None = None

    field_map = {

        "DOSAGE": "dosage", "STRENGTH": "dosage", "DOSE": "dosage",

        "FREQUENCY": "frequency", "ROUTE": "route",

    }

    for entity in sorted((e for e in entities if isinstance(e, dict)), key=_span_start):

        etype = str(entity.get("type", "")).strip().upper()

        value = str(entity.get("value", "")).strip()

        if not value:

            continue

        if etype in {"MEDICATION", "MEDICINE", "DRUG"}:

            current = medicines.setdefault(value.casefold(), _new_med(value))

        elif current is not None and etype in field_map:

            current[field_map[etype]] = value

    return [

        med for med in medicines.values()

        if med["medicine"].casefold() not in {"none", "null"}

    ]

def get_medication_details(result: dict[str, Any]) -> list[dict[str, Any]]:

    backend = result.get("medication_details")

    if isinstance(backend, list) and backend:

        merged = merge_duplicate_medications(backend)

        if merged:

            return merged

    entities = result.get("entities")

    if isinstance(entities, list):

        details = build_medications_from_entities(entities)

        if details:

            return details

    return [

        _new_med(m) for m in unique_strings(result.get("medicines", []))

        if m.casefold() not in {"none", "null"}

    ]

def medications_table(details: list[dict[str, Any]]) -> pd.DataFrame:

    rows = [

        {

            "Medicine": safe_text(item.get("medicine")),

            "Dosage": safe_text(item.get("dosage")),

            "Frequency": safe_text(item.get("frequency")),

            "Route": safe_text(item.get("route")),

        }

        for item in details

        if isinstance(item, dict)

        and str(item.get("medicine") or "").strip()

        and str(item.get("medicine")).strip().casefold() not in {"none", "null"}

    ]

    return pd.DataFrame(rows, columns=["Medicine", "Dosage", "Frequency", "Route"])

def medicine_symptom_links(result: dict[str, Any]) -> list[dict[str, str]]:

    relations = result.get("relations")

    if not isinstance(relations, list):

        return []

    rows, seen = [], set()

    for rel in relations:

        if not isinstance(rel, dict):

            continue

        medicine = rel.get("medication") or rel.get("medicine")

        symptom = rel.get("symptom")

        if not medicine or not symptom:

            continue

        assertion = rel.get("symptom_assertion")

        assertion = assertion if isinstance(assertion, dict) else {}

        status = (

            "Denied" if assertion.get("negated")

            else "Uncertain" if assertion.get("uncertain")

            else "Reported"

        )

        key = (str(medicine).casefold(), str(symptom).casefold(), status)

        if key in seen:

            continue

        seen.add(key)

        rows.append({"Medicine": str(medicine), "Symptom": str(symptom), "Status": status})

    return rows

def extract_text_from_file(uploaded_file: Any) -> str:

    """Extract and validate uploaded TXT, PDF, or DOCX content."""

    name = str(uploaded_file.name).lower()

    extension = Path(name).suffix

    max_file_size = 5 * 1024 * 1024

    max_text_length = 20000

    max_pdf_pages = 30

    if extension not in {".txt", ".pdf", ".docx"}:

        raise ValueError(

            "Unsupported file. Upload TXT, PDF or DOCX."

        )

    if uploaded_file.size == 0:

        raise ValueError("The uploaded file is empty.")

    if uploaded_file.size > max_file_size:

        raise ValueError(

            "File is too large. Maximum size is 5 MB."

        )

    data = uploaded_file.getvalue()

    try:

        if extension == ".txt":

            if b"\x00" in data:

                raise ValueError(

                    "Invalid TXT file. Binary content detected."

                )

            try:

                text = data.decode("utf-8-sig")

            except UnicodeDecodeError:

                raise ValueError(

                    "TXT file must use UTF-8 encoding."

                )

        elif extension == ".pdf":

            from pypdf import PdfReader

            if b"%PDF-" not in data[:1024]:

                raise ValueError("Invalid PDF file.")

            reader = PdfReader(BytesIO(data))

            if reader.is_encrypted:

                raise ValueError(

                    "Password-protected PDFs are not supported."

                )

            if len(reader.pages) > max_pdf_pages:

                raise ValueError(

                    "PDF exceeds the 30-page limit."

                )

            pages = [

                page.extract_text() or ""

                for page in reader.pages

            ]

            text = "\n\n".join(pages).strip()

            if not text:

                raise ValueError(

                    "No readable text found in PDF. "

                    "Scanned PDFs require OCR."

                )

        elif extension == ".docx":

            from docx import Document

            if not zipfile.is_zipfile(BytesIO(data)):

                raise ValueError("Invalid DOCX file.")

            # Basic protection against oversized ZIP content

            with zipfile.ZipFile(BytesIO(data)) as archive:

                total_size = sum(

                    info.file_size

                    for info in archive.infolist()

                )

                if total_size > 25 * 1024 * 1024:

                    raise ValueError(

                        "Document contains too much data."

                    )

            document = Document(BytesIO(data))

            paragraphs = [

                p.text for p in document.paragraphs

                if p.text.strip()

            ]

            # Include text inside Word tables

            for table in document.tables:

                for row in table.rows:

                    row_text = " ".join(

                        cell.text for cell in row.cells

                        if cell.text.strip()

                    )

                    if row_text.strip():

                        paragraphs.append(row_text)

            text = "\n".join(paragraphs).strip()

        text = text.strip()

        if not text:

            raise ValueError(

                "Document contains no readable narrative."

            )

        if len(text) > max_text_length:

            raise ValueError(

                "Document text exceeds 20,000 characters."

            )

        return text

    except ValueError:

        raise

    except ImportError:

        raise ValueError(

            "Required document-reading library is missing."

        )

    except Exception:

        raise ValueError(

            "Unable to read this document. "

            "Check that the file is valid and not corrupted."

        )

def format_timestamp(value: Any) -> str:

    if not value:

        return "—"

    try:

        text = str(value).strip()

        dt = datetime.fromisoformat(

            text.replace("Z", "+00:00")

        )

        # Assumes timezone-naive backend timestamps are UTC.

        if dt.tzinfo is None:

            dt = dt.replace(tzinfo=timezone.utc)

        ist = timezone(timedelta(hours=5, minutes=30))

        local_dt = dt.astimezone(ist)

        return local_dt.strftime("%Y-%m-%d %I:%M:%S %p")

    except (ValueError, TypeError):

        return str(value)

# SIDEBAR

# ============================================================

def apply_layout() -> None:

    """Keep spacing and colours consistent across Streamlit versions."""

    st.markdown(

        """

        <style>

        :root {

            --mr-primary: #0e6f8f;

            --mr-teal: #087e85;

            --mr-border: #e1e9ee;

            --mr-muted: #64748b;

            --mr-soft: #f3f8fa;

        }

        .block-container {

            max-width: 1230px;

            padding-top: 2.25rem;

            padding-bottom: 3rem;

            padding-left: 2.5rem;

            padding-right: 2.5rem;

        }

        [data-testid="stMain"] [data-testid="stVerticalBlock"] {

            gap: 0.95rem;

        }

        [data-testid="stHorizontalBlock"] { gap: 1.3rem; }

        h1, h2, h3 {

            color: #173445;

            letter-spacing: -0.025em;

            margin: 0 !important;

            padding: 0 !important;

        }

        h2 { font-size: 1.85rem !important; font-weight: 700 !important; }

        h3 { font-size: 1.24rem !important; font-weight: 650 !important; }

        [data-testid="stCaptionContainer"] {

            color: var(--mr-muted);

            line-height: 1.5;

        }

        [data-testid="stMarkdownContainer"] p { margin-bottom: 0; }

        [data-testid="stVerticalBlockBorderWrapper"] {

            border-radius: 15px;

            background: #fff;

        }

        [data-testid="stForm"] {

            border-color: var(--mr-border);

            border-radius: 15px;

            padding: 1.35rem 1.45rem 1.2rem;

        }

        [data-testid="stMetric"] {

            min-height: 100px;

            padding: 14px 16px;

            border: 1px solid var(--mr-border);

            border-radius: 12px;

            background: var(--mr-soft);

        }

        [data-testid="stMetricLabel"] { color: var(--mr-muted); }

        [data-testid="stMetricValue"] { color: #173445; }

        div[data-baseweb="input"] > div,

        div[data-baseweb="textarea"] > div,

        div[data-baseweb="select"] > div {

            border-radius: 9px;

        }

        [data-testid="stFileUploader"] section {

            border-radius: 12px;

            border-color: #cadde6;

            background: #f9fcfd;

        }

        [data-testid="stDataFrame"], [data-testid="stDataEditor"] {

            border-radius: 10px;

            overflow: hidden;

        }

        button[data-baseweb="tab"] {

            padding: 11px 18px;

            font-weight: 600;

        }

        button[data-baseweb="tab"][aria-selected="true"] {

            color: var(--mr-primary);

        }

        [data-baseweb="tab-panel"] { padding-top: 1rem; }

        .stButton > button, .stDownloadButton > button,

        [data-testid="stFormSubmitButton"] > button {

            border-radius: 10px;

            min-height: 41px;

            font-weight: 600;

        }

        hr {

            border-color: var(--mr-border) !important;

            margin: 0.85rem 0 !important;

        }

        [data-testid="stSidebar"] {

            background: #f3f6f9;

            border-right: 1px solid var(--mr-border);

        }

        [data-testid="stSidebarUserContent"] { padding-top: 2rem; }

        [data-testid="stSidebar"] [data-testid="stVerticalBlock"] { gap: 0.8rem; }

        [data-testid="stSidebar"] [role="radiogroup"] { gap: 0.45rem; }

        [data-testid="stSidebar"] [role="radiogroup"] label {

            border-radius: 9px;

            padding: 9px 10px;

        }

        [data-testid="stSidebar"] [role="radiogroup"] label:has(input:checked) {

            background: #e5f2f5;

        }

        .mr-brand {

            font-size: 1.42rem;

            font-weight: 750;

            color: #0e6f8f;

        }

        .mr-brand-sub {

            color: #64748b;

            font-size: 0.88rem;

            margin-top: 3px;

        }

        .mr-section-label {

            color: var(--mr-teal);

            font-size: 0.76rem;

            font-weight: 700;

            letter-spacing: 0.09em;

            text-transform: uppercase;

        }

        @media (max-width: 900px) {

            .block-container { padding: 1.6rem 1.15rem 2.5rem; }

            [data-testid="stHorizontalBlock"] { gap: 0.7rem; }

            [data-testid="stForm"] { padding: 1rem; }

            [data-testid="stMetric"] {

                padding: 11px;

                min-height: 88px;

            }

        }

        </style>

        """,

        unsafe_allow_html=True,

    )

def show_sidebar() -> str:

    with st.sidebar:

        st.markdown(

            "<div class='mr-brand'>✚ MediReview</div>"

            "<div class='mr-brand-sub'>Medicine safety review</div>",

            unsafe_allow_html=True,

        )

        st.markdown("<div class='mr-section-label'>Workspace</div>", unsafe_allow_html=True)

        view = st.radio("Navigate", VIEWS, key="active_view", label_visibility="collapsed")

        review_id = st.session_state.get("review_id")

        if review_id is not None:

            st.divider()

            st.caption("Current case")

            st.markdown(f"**Case {review_id}**")

            show_badge(st.session_state["reviewer_classification"])

            st.button("Close case", on_click=reset_current_case, use_container_width=True)

        st.divider()

        st.caption(

            "Medicine–symptom links are reported associations, "

            "not proof of causation."

        )

    return view

def validate_review_content(text: str) -> str | None:

    """Validate basic narrative quality and medical relevance.

    This is a conservative heuristic filter, not medical diagnosis or an

    exhaustive classifier. The original text is not altered so NLP evidence

    offsets and the saved narrative remain consistent.

    """

    import unicodedata

    if not isinstance(text, str):

        return "Please enter or upload a text narrative."

    if "\x00" in text:

        return "Invalid text content."

    if len(text) > 20_000:

        return "Review exceeds the maximum length of 20,000 characters."

    checked = unicodedata.normalize("NFKC", text)

    checked = re.sub(r"\s+", " ", checked).strip()

    if not checked:

        return "Please enter a medicine review."

    words = re.findall(r"[^\W\d_]+(?:['’][^\W\d_]+)?", checked, re.UNICODE)

    letters = sum(c.isalpha() for c in checked)

    digits = sum(c.isdigit() for c in checked)

    if not words or letters == 0:

        return "Review must contain descriptive text, not only numbers or symbols."

    if len(words) < 2:

        return "Please enter a descriptive medicine review, not a single word or code."

    tokens = {word.casefold() for word in words}

    word_list = [word.casefold() for word in words]

    # Detect clearly unrelated documents. Requiring multiple cues prevents

    # rejecting a patient report simply because it mentions work or education.

    unrelated_document_patterns = (

        r"\btechnical skills\b",

        r"\beducation\b",

        r"\b(?:academic|personal|professional) projects\b",

        r"\b(?:internships?|work experience|certifications?)\b",

        r"\b(?:bachelor of technology|b\.?\s?tech|bachelor of science)\b",

        r"\b(?:github\.com|linkedin\.com)\b",

        r"\b(?:computer science|data science|programming languages)\b",

        r"\b(?:machine learning|web development|professional summary)\b",

    )

    unrelated_hits = sum(

        bool(re.search(pattern, checked, re.IGNORECASE))

        for pattern in unrelated_document_patterns

    )

    if unrelated_hits >= 3:

        return (

            "This appears to be a resume or unrelated professional document. "

            "Please enter a medicine experience or upload a patient report."

        )

    # Strong medicine names and treatment terms; avoid treating ordinary

    # verbs such as 'took', 'started', or 'morning' as medical evidence alone.

    medicine_terms = {

        "medicine", "medicines", "medication", "medications", "drug", "drugs",

        "tablet", "tablets", "capsule", "capsules", "pills", "pill",

        "prescription", "prescribed", "antibiotic", "antibiotics",

        "inhaler", "injection", "injections", "syrup", "dosage", "dose",

        "treatment", "chemotherapy", "vaccination", "vaccine",

        "metformin", "paracetamol", "acetaminophen", "ibuprofen", "aspirin",

        "amoxicillin", "azithromycin", "doxycycline", "ciprofloxacin",

        "cefixime", "cephalexin", "metronidazole", "omeprazole",

        "pantoprazole", "cetirizine", "loratadine", "fexofenadine",

        "montelukast", "prednisone", "prednisolone", "dexamethasone",

        "salbutamol", "albuterol", "atorvastatin", "rosuvastatin",

        "amlodipine", "losartan", "telmisartan", "lisinopril",

        "metoprolol", "atenolol", "insulin", "glimepiride",

        "sitagliptin", "empagliflozin", "levothyroxine", "sertraline",

        "fluoxetine", "escitalopram", "gabapentin", "pregabalin",

        "tramadol", "diclofenac", "naproxen", "warfarin", "apixaban",

        "clopidogrel", "ondansetron", "loperamide",

    }

    symptom_terms = {

        "nausea", "nauseous", "vomit", "vomiting", "headache", "headaches",

        "dizziness", "dizzy", "vertigo", "fever", "chills", "fatigue",

        "tiredness", "weakness", "drowsiness", "drowsy", "insomnia",

        "rash", "rashes", "itching", "hives", "swelling", "redness",

        "pain", "painful", "cramps", "bloating", "diarrhea", "diarrhoea",

        "constipation", "indigestion", "heartburn", "stomachache",

        "cough", "coughing", "wheezing", "breathlessness", "palpitations",

        "sweating", "tremor", "tremors", "numbness", "tingling",

        "blurred", "dehydration", "anxiety", "restlessness", "confusion",

        "fainting", "seizure", "bleeding", "bruising", "burning",

        "allergy", "allergic", "anaphylaxis", "sleeplessness", "unwell",

    }

    has_medicine = bool(tokens & medicine_terms)

    has_symptom = bool(tokens & symptom_terms)

    has_dose = bool(re.search(

        r"\b\d+(?:[.,]\d+)?\s*(?:mg|mcg|µg|g|ml|iu)\b",

        checked, re.IGNORECASE

    ))

    has_event_phrase = bool(re.search(

        r"\b(?:side effects?|adverse (?:events?|reactions?)|"

        r"no (?:symptoms?|reactions?)|without (?:any )?(?:side effects?|symptoms?)|"

        r"medication history|medicine review|patient report)\b",

        checked, re.IGNORECASE

    ))

    # Common junk patterns. Check before sending text to the backend.

    if len(word_list) >= 4 and len(set(word_list)) <= 2:

        return "Please enter a narrative rather than repeated words."

    if re.search(r"([A-Za-z0-9])\1{6,}", checked) and not (

        has_medicine or has_symptom or has_event_phrase

    ):

        return "Review appears to contain repeated or nonsensical characters."

    if digits >= 8 and digits > letters and not (

        has_medicine or has_symptom or has_event_phrase or has_dose

    ):

        return "Review contains too much numeric text to be a narrative."

    if not (has_medicine or has_symptom or has_dose or has_event_phrase):

        return (

            "No clear medicine or symptom context was found. "

            "Please describe a patient medicine-related experience."

        )

    # A clinical report can legitimately have no named medicine/symptom and

    # should still be allowed to reach 'Insufficient Information'.

    return None

def show_new_case() -> None:

    st.header("New case", anchor=False)

    st.caption("Enter a medicine experience narrative or upload a document for review.")

    with st.form("analysis_form"):

        st.markdown("**Case references**")

        c1, c2 = st.columns(2)

        patient_reference = c1.text_input(

            "Patient reference", key="patient_reference_input", placeholder="Optional patient ID"

        )

        report_reference = c2.text_input(

            "Report reference", key="report_reference_input", placeholder="Optional report ID"

        )

        st.markdown("**Review content**")

        write_tab, upload_tab = st.tabs(["Write narrative", "Upload document"])

        with write_tab:

            narrative = st.text_area(

                "Medicine experience",

                height=220,

                key="manual_review",

                placeholder=(

                    "Patient took Metformin 500 mg twice daily and later reported "

                    "nausea, but denied rash."

                ),

            )

        with upload_tab:

            upload = st.file_uploader(

                "Upload report (TXT, PDF or DOCX)",

                type=["txt", "pdf", "docx"],

                key="uploaded_report",

            )

        _, button_col = st.columns([3, 1])

        with button_col:

            submitted = st.form_submit_button(

                "Analyze case", type="primary", use_container_width=True

            )

    if not submitted:

        return

    if (narrative or "").strip() and upload is not None:

        st.error("Use either the typed narrative or an uploaded document, not both.")

        return

    text, source_name = (narrative or "").strip(), "Typed narrative"

    if not text and upload is not None:

        try:

            text = extract_text_from_file(upload)

        except ValueError as exc:

            st.error(str(exc))

            return

        source_name = upload.name

    # Validate both manually entered text and extracted PDF/TXT/DOCX text.

    # Do this BEFORE create_review() so rejected input cannot create a case.

    error = validate_review_content(text)

    if error:

        st.error(error)

        return

    try:

        with st.spinner("Analyzing and saving..."):

            response = create_review(

                text,

                patient_reference=patient_reference.strip() or None,

                report_reference=report_reference.strip() or None,

                source_name=source_name,

            )

        result = response.get("analysis_result")

        if not isinstance(result, dict):

            raise ValueError("Response has no valid analysis result.")

        result["review_id"] = response.get("id")

        result["patient_reference"] = response.get("patient_reference")

        result["report_reference"] = response.get("report_reference")

    except Exception as exc:

        st.error(f"Unable to analyze case: {exc}")

        return

    st.session_state.update(

        analysis_result=result,

        submitted_text=text,

        source_name=source_name,

        review_id=response.get("id"),

        reviewer_classification="Needs Review",

        reviewer_notes="",

        history_loaded=False,

        goto="Current case",

    )

    st.rerun()

TEAL = colors.HexColor("#087C88")
INK = colors.HexColor("#203744")
MUTED = colors.HexColor("#647A88")
BORDER = colors.HexColor("#DDE7EB")
LIGHT = colors.HexColor("#F2F8F9")
WHITE = colors.white
IST = timezone(timedelta(hours=5, minutes=30))


def _get_fonts() -> tuple[str, str]:
    """Use a widely available Unicode font when possible (including Windows)."""
    regular_paths = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
        Path("C:/Windows/Fonts/segoeui.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    bold_paths = [
        Path("C:/Windows/Fonts/arialbd.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf"),
        Path("C:/Windows/Fonts/segoeuib.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ]
    regular = next((p for p in regular_paths if p.is_file()), None)
    bold = next((p for p in bold_paths if p.is_file()), None)
    if regular is None:
        return "Helvetica", "Helvetica-Bold"
    if "MRReportRegular" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("MRReportRegular", str(regular)))
    if bold is None:
        return "MRReportRegular", "MRReportRegular"
    if "MRReportBold" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("MRReportBold", str(bold)))
    return "MRReportRegular", "MRReportBold"


def _as_text(value: Any, default: str = "Not provided") -> str:
    if value is None or value == "":
        return default
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def _xml(value: Any) -> str:
    """Escape dynamic user text so it cannot break ReportLab Paragraph markup."""
    text = _as_text(value)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", text)
    return escape(text).replace("\n", "<br/>")


def _unique(values: Any) -> list[str]:
    if not isinstance(values, (list, tuple)):
        return []
    seen = set()
    result = []
    for value in values:
        text = _as_text(value, "").strip()
        if text and text.casefold() not in seen:
            result.append(text)
            seen.add(text.casefold())
    return result


def create_pdf_report(
    result: dict[str, Any],
    *,
    case_id: Any,
    patient_reference: Any,
    report_reference: Any,
    source_name: Any,
    original_narrative: str,
    medication_details: list[dict[str, Any]],
    reviewer_classification: str,
    reviewer_notes: str,
) -> bytes:
    """Generate PDF in memory: nothing is written to a server-side file."""
    if not isinstance(result, dict):
        raise ValueError("Missing NLP analysis to include in the PDF.")

    regular, bold = _get_fonts()
    styles = {
        "title": ParagraphStyle(
            "MRTitle", fontName=bold, fontSize=19, leading=24,
            textColor=INK, spaceAfter=3,
        ),
        "subtitle": ParagraphStyle(
            "MRSubtitle", fontName=regular, fontSize=9, leading=13,
            textColor=MUTED, spaceAfter=12,
        ),
        "section": ParagraphStyle(
            "MRSection", fontName=bold, fontSize=11, leading=15,
            textColor=TEAL, spaceBefore=14, spaceAfter=7,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "MRBody", fontName=regular, fontSize=8.6, leading=13,
            textColor=INK, spaceAfter=5, splitLongWords=True,
        ),
        "small": ParagraphStyle(
            "MRSmall", fontName=regular, fontSize=7.5, leading=11,
            textColor=INK, spaceAfter=3, splitLongWords=True,
        ),
        "table_head": ParagraphStyle(
            "MRTableHead", fontName=bold, fontSize=7.8, leading=11,
            textColor=WHITE,
        ),
        "notice": ParagraphStyle(
            "MRNotice", fontName=regular, fontSize=7.5, leading=11,
            textColor=MUTED, spaceAfter=8,
        ),
    }

    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=19 * mm, bottomMargin=19 * mm,
        title=f"MediReview Case {_as_text(case_id, 'Report')}",
        author="MediReview",
    )
    story: list[Any] = []

    def para(value: Any, style: str = "body") -> Paragraph:
        return Paragraph(_xml(value), styles[style])

    def section(title: str) -> None:
        story.append(Paragraph(_xml(title), styles["section"]))

    def matrix(headers: list[str], rows: list[list[Any]], widths: list[float] | None = None) -> None:
        if not rows:
            story.append(para("No information recorded."))
            return
        content: list[list[Any]] = [
            [Paragraph(_xml(h), styles["table_head"]) for h in headers]
        ]
        content.extend([[para(v, "small") for v in row] for row in rows])
        table = LongTable(
            content, colWidths=widths, repeatRows=1, hAlign="LEFT",
            splitByRow=1,
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), TEAL),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LIGHT]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, TEAL),
            ("LINEBELOW", (0, 1), (-1, -1), 0.3, BORDER),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        story.append(table)
        story.append(Spacer(1, 3 * mm))

    story.append(Paragraph("MediReview", styles["title"]))
    story.append(Paragraph(
        "MEDICINE SAFETY REVIEW REPORT  |  NLP EVIDENCE &amp; REVIEWER ASSESSMENT",
        styles["subtitle"],
    ))
    story.append(HRFlowable(width="100%", color=BORDER, thickness=0.8))

    section("1. Case details")
    matrix(["Field", "Value"], [
        ["Case ID", _as_text(case_id)],
        ["Patient reference", _as_text(patient_reference)],
        ["Report reference", _as_text(report_reference)],
        ["Source", _as_text(source_name)],
        ["Report exported (IST)", datetime.now(IST).strftime("%d %b %Y, %I:%M %p")],
    ], [46 * mm, 128 * mm])

    section("2. Automated and reviewer assessments")
    matrix(["Assessment", "Result"], [
        ["NLP classification", _as_text(result.get("narrative_classification"), "Not classified")],
        ["Reviewer classification", _as_text(reviewer_classification, "Needs Review")],
    ], [46 * mm, 128 * mm])
    story.append(Paragraph("<b>Reviewer notes</b>", styles["body"]))
    note_lines = _as_text(reviewer_notes, "No notes recorded").splitlines()
    for note_line in note_lines or ["No notes recorded"]:
        for offset in range(0, len(note_line), 850):
            story.append(para(note_line[offset:offset + 850]))
        if not note_line.strip():
            story.append(Spacer(1, 2 * mm))

    section("3. Identified medicines")
    medications = [m for m in medication_details if isinstance(m, dict)]
    med_rows = []
    for med in medications:
        med_rows.append([
            med.get("medicine") or med.get("Medicine") or med.get("name"),
            med.get("dosage") or med.get("Dosage"),
            med.get("frequency") or med.get("Frequency"),
            med.get("route") or med.get("Route"),
        ])
    matrix(["Medicine", "Dosage", "Frequency", "Route"], med_rows,
           [51 * mm, 37 * mm, 47 * mm, 39 * mm])

    section("4. Reported and denied symptoms")
    reported = _unique(result.get("reported_symptoms"))
    denied = _unique(result.get("negated_symptoms"))
    uncertain = _unique(result.get("uncertain_symptoms"))
    matrix(["Symptom status", "Extracted symptoms"], [
        ["Reported", ", ".join(reported) if reported else "None identified"],
        ["Denied / negated", ", ".join(denied) if denied else "None identified"],
        ["Uncertain", ", ".join(uncertain) if uncertain else "None identified"],
    ], [46 * mm, 128 * mm])

    entities = result.get("entities")
    entities = entities if isinstance(entities, list) else []
    section("5. Timeline evidence")
    timeline = []
    for entity in entities:
        if isinstance(entity, dict) and str(entity.get("type", "")).upper() in {"DATE", "DURATION", "TIME"}:
            timeline.append([entity.get("type"), entity.get("value")])
    matrix(["Type", "Evidence text"], timeline, [46 * mm, 128 * mm])

    section("6. Medicine-symptom associations")
    relations = result.get("relations")
    relations = relations if isinstance(relations, list) else []
    rel_rows = []
    for relation in relations:
        if not isinstance(relation, dict):
            continue
        assertion = relation.get("symptom_assertion")
        assertion = assertion if isinstance(assertion, dict) else {}
        status = (
            "Denied" if assertion.get("negated")
            else "Uncertain" if assertion.get("uncertain")
            else "Reported"
        )
        rel_rows.append([
            relation.get("medication") or relation.get("medicine"),
            relation.get("symptom"), status,
        ])
    matrix(["Medicine", "Symptom", "Evidence status"], rel_rows,
           [67 * mm, 67 * mm, 40 * mm])
    story.append(para(
        "These links are text-based reported associations and do not establish causality.",
        "notice",
    ))

    section("7. Extracted entities and evidence spans")
    entity_rows = []
    for entity in entities:
        if not isinstance(entity, dict):
            continue
        span = entity.get("span")
        if isinstance(span, dict):
            start = span.get("char_start", span.get("start"))
            end = span.get("char_end", span.get("end"))
            location = f"{start}-{end}" if start is not None and end is not None else "-"
        else:
            location = "-"
        assertion = entity.get("assertion")
        if not isinstance(assertion, dict):
            assertion = {}
        flags = [
            name for name, match in (
                ("negated", assertion.get("negated", entity.get("negated"))),
                ("uncertain", assertion.get("uncertain", entity.get("uncertain"))),
                ("historical", assertion.get("historical", entity.get("historical"))),
            ) if match
        ]
        entity_rows.append([
            entity.get("type"), entity.get("value"),
            location, ", ".join(flags) if flags else "-",
        ])
    matrix(["Entity type", "Exact text", "Offsets", "Assertion"], entity_rows,
           [40 * mm, 78 * mm, 24 * mm, 32 * mm])

    section("8. Original patient narrative")
    raw_narrative = _as_text(original_narrative, "No original narrative available")
    # Split long narratives into printable paragraphs to avoid a giant unbreakable block.
    for original_line in raw_narrative.splitlines() or [raw_narrative]:
        if not original_line.strip():
            story.append(Spacer(1, 2 * mm))
            continue
        for offset in range(0, len(original_line), 850):
            story.append(para(original_line[offset:offset + 850]))

    def decorate(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        w, h = A4
        canvas.setStrokeColor(BORDER)
        canvas.setLineWidth(0.45)
        canvas.line(18 * mm, 14 * mm, w - 18 * mm, 14 * mm)
        canvas.setFont(regular, 7)
        canvas.setFillColor(MUTED)
        canvas.drawString(18 * mm, 10 * mm, "MediReview  |  Medicine safety review evidence")
        canvas.drawRightString(w - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return output.getvalue()



def show_export(result: dict[str, Any]) -> None:
    """Download a complete case PDF with NLP evidence and saved reviewer assessment."""
    try:
        pdf_bytes = create_pdf_report(
            result=result,
            case_id=result.get("review_id", st.session_state.get("review_id")),
            patient_reference=result.get("patient_reference"),
            report_reference=result.get("report_reference"),
            source_name=st.session_state.get("source_name"),
            original_narrative=st.session_state.get("submitted_text", ""),
            medication_details=get_medication_details(result),
            reviewer_classification=st.session_state.get(
                "reviewer_classification", "Needs Review"
            ),
            reviewer_notes=st.session_state.get("reviewer_notes", ""),
        )
        case_id = safe_text(result.get("review_id"), "new")
        st.download_button(
            "Save PDF",
            data=pdf_bytes,
            file_name=f"medireview_case_{case_id}.pdf",
            mime="application/pdf",
            use_container_width=True,
            key=f"download_case_pdf_{case_id}",
        )
    except Exception as exc:
        st.error(f"Unable to prepare PDF report: {exc}")


def show_case_header(result: dict[str, Any]) -> None:

    st.header(f"Case {safe_text(result.get('review_id'))}", anchor=False)

    st.caption(

        f"Patient {safe_text(result.get('patient_reference'))}  ·  "

        f"Report {safe_text(result.get('report_reference'))}  ·  "

        f"{safe_text(st.session_state.get('source_name'))}"

    )

    with st.container(border=True):

        status_left, status_right = st.columns(2)

        with status_left:

            st.caption("Automated assessment")

            show_badge(result.get("narrative_classification"))

        with status_right:

            st.caption("Reviewer assessment")

            show_badge(st.session_state["reviewer_classification"])

    med_col, reported_col, denied_col = st.columns(3)

    med_col.metric("Medicines", len(get_medication_details(result)))

    reported_col.metric("Symptoms reported", len(unique_strings(result.get("reported_symptoms", []))))

    denied_col.metric("Symptoms denied", len(unique_strings(result.get("negated_symptoms", []))))

def _bullets(items: list[str], empty: str) -> None:

    if items:

        st.markdown("\n".join(f"- {item}" for item in items))

    else:

        st.caption(empty)

def show_evidence(result: dict[str, Any]) -> None:

    # Display the original NLP findings.

    meds = get_medication_details(result)

    with st.container(border=True):

        st.markdown("**Original NLP medicine information**")


        if meds:

            st.dataframe(medications_table(meds), hide_index=True, use_container_width=True)

        else:

            st.caption("No medicines identified.")

    reported = unique_strings(result.get("reported_symptoms", []))

    denied = unique_strings(result.get("negated_symptoms", []))

    left, right = st.columns(2)

    with left:

        with st.container(border=True):

            st.markdown("**Symptoms reported**")

            _bullets(reported, "None identified.")

    with right:

        with st.container(border=True):

            st.markdown("**Symptoms denied**")

            _bullets(denied, "None identified.")

    timing = [*get_entity_values(result, "DATE"), *get_entity_values(result, "DURATION")]

    if timing:

        with st.container(border=True):

            st.markdown("**Timeline**")

            st.write(" · ".join(timing))

    links = medicine_symptom_links(result)

    if links:

        with st.container(border=True):

            st.markdown("**Medicine–symptom mentions**")

            st.dataframe(pd.DataFrame(links), hide_index=True, use_container_width=True)

            st.caption("Reported associations only; causation is not inferred.")

def show_reviewer_assessment() -> None:

    review_id = st.session_state.get("review_id")

    key = str(review_id)

    current = st.session_state["reviewer_classification"]

    index = REVIEWER_OPTIONS.index(current) if current in REVIEWER_OPTIONS else 0

    with st.container(border=True):

        st.markdown("**Assessment**")

        classification = st.selectbox(

            "Classification", REVIEWER_OPTIONS, index=index, key=f"reviewer_classification_select_{key}"

        )

        notes = st.text_area(

            "Notes",

            value=st.session_state["reviewer_notes"],

            height=160,

            key=f"reviewer_notes_field_{key}",

        )

        if st.button("Save assessment", type="primary", use_container_width=True, key=f"save_assessment_{key}"):

            try:

                with st.spinner("Saving..."):

                    response = update_review_assessment(int(review_id), classification, notes.strip())

                st.session_state["reviewer_classification"] = response.get(

                    "reviewer_classification", classification

                )

                st.session_state["reviewer_notes"] = response.get("reviewer_notes", notes)

                st.session_state["history_loaded"] = False
                st.rerun()


            except Exception as exc:

                st.error(f"Unable to save: {exc}")

def show_current_case() -> None:

    result = st.session_state.get("analysis_result")

    if not isinstance(result, dict):

        st.header("Current case", anchor=False)

        st.info("No case is open. Create a new case or open one from History.")

        return

    show_case_header(result)

    evidence_tab, review_tab, narrative_tab = st.tabs(

        ["Evidence", "Reviewer workspace", "Narrative"]

    )

    with evidence_tab:

        show_evidence(result)

    with review_tab:

        show_reviewer_assessment()

    with narrative_tab:

        with st.container(border=True):

            st.markdown("**Original narrative**")

            st.write(st.session_state.get("submitted_text", "") or "—")

    # Place the report action after the evidence/reviewer/narrative tabs.
    # The PDF always reflects the last successfully saved reviewer assessment.
    st.divider()
    _, pdf_col = st.columns([4, 1])
    with pdf_col:
        show_export(result)


def show_history() -> None:

    st.header("History", anchor=False)

    st.caption("Search saved reviews and open a case to inspect its evidence.")

    if st.session_state.get("history_error"):

        st.error(st.session_state["history_error"])

    with st.form("history_filter_form"):

        st.markdown("**Filter cases**")

        c1, c2, c3 = st.columns(3)

        search = c1.text_input("Search", key="history_search", placeholder="Review text or case ID")

        medicine = c2.text_input("Medicine", key="history_medicine", placeholder="Medicine name")

        classification = c3.selectbox("Classification", FILTER_OPTIONS, key="history_classification")

        c4, c5 = st.columns(2)

        patient_reference = c4.text_input("Patient reference", key="history_patient_reference")

        report_reference = c5.text_input("Report reference", key="history_report_reference")

        _, search_col = st.columns([4, 1])

        with search_col:

            submitted = st.form_submit_button("Search cases", type="primary", use_container_width=True)

    if submitted or not st.session_state["history_loaded"]:

        try:

            with st.spinner("Loading..."):

                if submitted:

                    reviews = get_reviews(

                        classification=None if classification == "All" else classification,

                        patient_reference=patient_reference.strip() or None,

                        report_reference=report_reference.strip() or None,

                        medicine=medicine.strip() or None,

                        search=search.strip() or None,

                    )

                else:

                    reviews = get_reviews()

            st.session_state["history_reviews"] = reviews if isinstance(reviews, list) else []

            st.session_state["history_loaded"] = True

        except Exception as exc:

            st.error(f"Unable to load history: {exc}")

            return

    reviews = [r for r in st.session_state["history_reviews"] if isinstance(r, dict)]

    if not reviews:

        st.info("No cases match.")

        return

    # Selection is from the previous run so the action can sit above the table.

    state = st.session_state.get("history_table") or {}

    rows = (state.get("selection") or {}).get("rows") or []

    selected_id = None

    if rows and rows[0] < len(reviews):

        try:

            selected_id = int(reviews[rows[0]].get("id"))

        except (TypeError, ValueError):

            selected_id = None

    count_col, open_col = st.columns([4, 1], vertical_alignment="center")

    count_col.caption(f"{len(reviews)} case(s) · select a row to open it")

    open_col.button(

        "Open case", type="primary", disabled=selected_id is None,

        on_click=open_saved_case, args=(selected_id,), use_container_width=True,

    )

    table = pd.DataFrame([

        {

            "Case": r.get("id"),

            "Patient": safe_text(r.get("patient_reference")),

            "Report": safe_text(r.get("report_reference")),

            "Automated": safe_text(r.get("narrative_classification"), "Not classified"),

            "Reviewer": safe_text(r.get("reviewer_classification"), "Needs Review"),

            "Created": format_timestamp(r.get("created_at")),

        }

        for r in reviews

    ])

    st.dataframe(

        table, hide_index=True, on_select="rerun", selection_mode="single-row",

        key="history_table", use_container_width=True,

    )

def main() -> None:

    initialize_state()

    apply_layout()

    # Apply a pending view change before the navigation widget is created.

    if st.session_state["goto"]:

        st.session_state["active_view"] = st.session_state["goto"]

        st.session_state["goto"] = None

    view = show_sidebar()

    if view == "New case":

        show_new_case()

    elif view == "Current case":

        show_current_case()

    else:

        show_history()

if __name__ == "__main__":

    main()
